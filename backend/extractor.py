import base64
from pathlib import Path
from typing import Tuple


SUPPORTED = {'.pdf', '.txt', '.md', '.docx', '.pptx', '.jpg', '.jpeg', '.png'}
IMAGE_EXTS = {'.jpg', '.jpeg', '.png'}
IMAGE_MIME = {'.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.png': 'image/png'}


def extract(file_bytes: bytes, filename: str) -> Tuple[str, bool, str]:
    """
    Extract text from a file.
    Returns (content, is_image, mime_type)
    - is_image=True  → content is base64, send to vision API
    - is_image=False → content is plain text
    """
    ext = Path(filename).suffix.lower()

    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported type '{ext}'. Supported: {', '.join(sorted(SUPPORTED))}")

    if ext in IMAGE_EXTS:
        b64 = base64.b64encode(file_bytes).decode()
        return b64, True, IMAGE_MIME.get(ext, 'image/jpeg')

    if ext == '.pdf':
        return _pdf(file_bytes), False, ''

    if ext in ('.txt', '.md'):
        return file_bytes.decode('utf-8', errors='ignore'), False, ''

    if ext == '.docx':
        return _docx(file_bytes), False, ''

    if ext == '.pptx':
        return _pptx(file_bytes), False, ''

    raise ValueError(f"Unhandled extension: {ext}")


# ── Extractors ────────────────────────────────────────────────────────────────

def _pdf(data: bytes) -> str:
    import io
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or '' for page in reader.pages]
    return '\n\n'.join(p for p in pages if p.strip())


def _docx(data: bytes) -> str:
    import io
    from docx import Document
    doc = Document(io.BytesIO(data))
    return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())


def _pptx(data: bytes) -> str:
    import io
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                lines.append(shape.text.strip())
    return '\n'.join(lines)
